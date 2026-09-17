"""Build a spreadsheet for checking the ranks the change report flags.

    python -m pipeline.build_rank_sheet                       # -> data/review/rank_review.xlsx
    python -m pipeline.import_rank_review data/review/rank_review.xlsx

One row per flagged appointment: everyone held as assistant or associate for
RANK_REVIEW_YEARS or more, or (for the inherited roster, which has no rank
start dates) an assistant that long past the PhD. A reviewer opens the Scholar
profile and the program website, picks what the person's rank actually is from
a dropdown, notes where they saw it, and the importer applies it.

Bands, longest-held first: A is 14+ years, B is 10-13, C is 7-9.
"""

from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import BUILD_DIR, ROSTER_DIR
from .build_review_sheet import ASK_FILL, BAND_FILL, GROUP_FILL, HAIR, HEAD_FILL, THIN
from .detect_changes import RANK_REVIEW_YEARS, collect_rank_review, latest_two, rows
from .match_openalex import REVIEW_DIR

SCHOLAR_URL = "https://scholar.google.com/citations?user={}"
OUTCOMES = ["assistant", "associate", "full", "emeritus", "adjunct or lecturer", "left the program", "unclear"]
HEADERS = ["Band", "Name", "Program", "Rank on file", "PhD year", "Years", "Basis", "Scholar profile says",
           "Citations", "Pct in rank", "Scholar profile", "Program website",
           "Actual rank", "Since (year)", "Where you saw it", "Reviewed by", "Date", "Notes",
           "person_id", "affiliation_id"]
BANDS = {"A": "14 or more years. Almost certainly a stale rank, or not a tenure-track post at all.",
         "B": "10 to 13 years.", "C": f"{RANK_REVIEW_YEARS} to 9 years."}


def band_of(years: int) -> str:
    return "A" if years >= 14 else "B" if years >= 10 else "C"


def load(db_path: Path):
    from datetime import date
    people = {p["person_id"]: p for p in rows(ROSTER_DIR / "person.csv")}
    depts = {d["department_id"]: d for d in rows(ROSTER_DIR / "department.csv")}
    current = {a["person_id"]: a for a in rows(ROSTER_DIR / "affiliation.csv")
               if not a["end_date"] and a["is_primary"] == "1"}
    gs_latest, _ = latest_two("google_scholar")
    gs_text = {r["person_id"]: (json.loads(r["raw_json"] or "{}").get("affiliation") or "") for r in gs_latest}
    metrics = {}
    con = sqlite3.connect(db_path)
    for pid, pct, cites in con.execute("SELECT person_id, pct_citations_rank, total_citations FROM v_person_percentiles"):
        metrics[str(pid)] = (pct, cites)
    con.close()
    return collect_rank_review(people, depts, current, gs_text, metrics, date.today())


def write_sheet(path: Path, review: list[dict]) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "C2"
    col = {h: i + 1 for i, h in enumerate(HEADERS)}

    dv = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES) + '"', allow_blank=True, showDropDown=False)
    dv.promptTitle = "Actual rank"
    dv.prompt = "What the program website or profile shows today. Leave blank to decide later."
    ws.add_data_validation(dv)

    for i, r in enumerate(review, start=2):
        band = band_of(r["years"])
        ws.cell(i, col["Band"], band).fill = BAND_FILL
        ws.cell(i, col["Name"], r["display_name"])
        ws.cell(i, col["Program"], r["department"])
        ws.cell(i, col["Rank on file"], r["rank"])
        ws.cell(i, col["PhD year"], int(r["phd_year"]) if r["phd_year"] else None)
        ws.cell(i, col["Years"], r["years"])
        ws.cell(i, col["Basis"], r["basis"])
        ws.cell(i, col["Scholar profile says"], r["profile_affiliation"] or "(no Scholar profile)")
        ws.cell(i, col["Citations"], int(r["total_citations"]) if r["total_citations"] != "" else None)
        ws.cell(i, col["Pct in rank"], round(float(r["pct_citations_rank"]), 2) if r["pct_citations_rank"] != "" else None)
        if r["google_scholar_id"]:
            c = ws.cell(i, col["Scholar profile"], "open profile")
            c.hyperlink = SCHOLAR_URL.format(r["google_scholar_id"]); c.font = Font(color="0563C1", underline="single")
        if r["department_url"]:
            c = ws.cell(i, col["Program website"], "open site")
            c.hyperlink = r["department_url"]; c.font = Font(color="0563C1", underline="single")
        ws.cell(i, col["person_id"], r["person_id"])
        ws.cell(i, col["affiliation_id"], r["affiliation_id"])
        dv.add(ws.cell(i, col["Actual rank"]))
        for h in ("Actual rank", "Since (year)", "Where you saw it", "Reviewed by", "Date", "Notes"):
            ws.cell(i, col[h]).fill = ASK_FILL
        for c in ws[i]:
            c.border = Border(bottom=HAIR)
            if c.column in (col["Basis"], col["Scholar profile says"], col["Notes"], col["Where you saw it"]):
                c.alignment = Alignment(wrap_text=True, vertical="top")
        if i > 2 and band_of(review[i - 3]["years"]) != band:
            for c in ws[i - 1]:
                c.border = Border(bottom=THIN)

    widths = {"Band": 6, "Name": 22, "Program": 26, "Rank on file": 11, "PhD year": 8, "Years": 6, "Basis": 34,
              "Scholar profile says": 34, "Citations": 10, "Pct in rank": 8, "Scholar profile": 13,
              "Program website": 13, "Actual rank": 17, "Since (year)": 10, "Where you saw it": 30,
              "Reviewed by": 13, "Date": 11, "Notes": 30, "person_id": 9, "affiliation_id": 11}
    for h, w in widths.items():
        ws.column_dimensions[get_column_letter(col[h])].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(review) + 1}"

    doc = wb.create_sheet("How to use this")
    lines = [
        ("What this is", True),
        (f"{len(review)} faculty we hold as assistant or associate professor who have been at that rank for "
         f"{RANK_REVIEW_YEARS} years or more, or look that way. The site updates a rank only when a person's own "
         "Google Scholar profile states one, and many profiles never do, so promotions get missed. Most of these "
         "ranks came from a hand-kept spreadsheet and some are years out of date.", False),
        ("", False),
        ("What to do", True),
        ("For each row, open the Scholar profile and the program website (the two link columns) and find the "
         "person's current title. Pick it in 'Actual rank'. If the page gives a year they were promoted, put it in "
         "'Since (year)'. Paste the page address in 'Where you saw it'. Leave a row blank to come back to it.", False),
        ("", False),
        ("The options", True),
        ("assistant / associate / full   their tenure-line rank today. Picking the same rank that is on file "
         "records that it was checked, so the person is not raised again for seven years.", False),
        ("emeritus                        retired with the title. They come off the site.", False),
        ("adjunct or lecturer             not a tenure-line post (lecturer, instructor, professor of practice, "
         "adjunct, research or clinical). They come off the site; the roster tracks tenure-line faculty only.", False),
        ("left the program                no longer at this program. Note where they went if you can see it.", False),
        ("unclear                         you looked and could not tell. Say why in Notes.", False),
        ("", False),
        ("Reading the evidence columns", True),
        ("Years / Basis          how long we think they have held the rank, and why we think so. 'rank start not on "
         "record' means we are going by years since the PhD.", False),
        ("Scholar profile says   the affiliation line on their profile, exactly as written. If it says "
         "'Associate Professor' and Rank on file says 'assistant', that is your answer.", False),
        ("Pct in rank            their citation count as a percentile among the people we hold at the same rank. "
         "0.98 for an assistant professor is a strong hint the rank is stale.", False),
        ("", False),
        ("Bands", True),
    ] + [(f"{k}  {v}", False) for k, v in BANDS.items()] + [
        ("", False),
        ("When you are done", True),
        ("python -m pipeline.import_rank_review data/review/rank_review.xlsx", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = doc.cell(i, 1, text)
        c.font = Font(bold=bold, size=12 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    doc.column_dimensions["A"].width = 112

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    tmp.replace(path)
    return len(review)


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db", type=Path, default=BUILD_DIR / "citations.sqlite")
    ap.add_argument("--out", type=Path, default=REVIEW_DIR / "rank_review.xlsx")
    args = ap.parse_args(argv)
    review = load(args.db)
    if not review:
        raise SystemExit("nobody is due a rank check")
    n = write_sheet(args.out, review)
    bands = {b: sum(1 for r in review if band_of(r["years"]) == b) for b in "ABC"}
    print(f"wrote {args.out}: {n} people  " + "  ".join(f"band {b}: {k}" for b, k in bands.items()))


if __name__ == "__main__":
    main()
