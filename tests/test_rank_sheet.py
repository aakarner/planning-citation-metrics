"""The rank-review spreadsheet and its importer.

The sheet is written by one script and read by another, so the columns the
importer looks up by name are a contract. apply() is pure so every outcome a
reviewer can pick is checked against what it does to the roster.
"""

from openpyxl import load_workbook

from pipeline.build_rank_sheet import HEADERS, OUTCOMES, band_of, write_sheet
from pipeline.import_rank_review import apply, read_decisions, since_date


def row(**kw):
    base = {"person_id": "1", "affiliation_id": "10", "display_name": "Ada Lovelace", "department": "Somewhere",
            "department_url": "https://planning.somewhere.edu/", "rank": "assistant", "years": 12,
            "basis": "assistant 12 years after the 2014 PhD; rank start not on record", "phd_year": 2014,
            "rank_start": "", "profile_affiliation": "Somewhere University", "pct_citations_rank": 0.97,
            "total_citations": 4000, "google_scholar_id": "abcdefghijkl"}
    return {**base, **kw}


def aff(**kw):
    base = {"affiliation_id": "10", "person_id": "1", "department_id": "5", "rank": "assistant",
            "appointment_type": "regular", "is_primary": "1", "start_date": "", "end_date": "", "source": "workbook-2026"}
    return {**base, **kw}


def decision(**kw):
    base = {"affiliation_id": "10", "person_id": "1", "on_file": "assistant", "decision": "associate",
            "since": None, "url": "https://somewhere.edu/people/ada", "by": "RA", "date": "2026-09-17", "notes": ""}
    return {**base, **kw}


def test_bands_run_longest_first():
    assert band_of(14) == "A" and band_of(20) == "A"
    assert band_of(10) == "B" and band_of(13) == "B"
    assert band_of(7) == "C" and band_of(9) == "C"


def test_the_sheet_carries_every_column_the_importer_needs(tmp_path):
    out = tmp_path / "r.xlsx"
    write_sheet(out, [row()])
    ws = load_workbook(out)["Review"]
    hdr = [c.value for c in ws[1]]
    for needed in ("Actual rank", "Since (year)", "Where you saw it", "Reviewed by", "Date", "Notes",
                   "person_id", "affiliation_id", "Rank on file"):
        assert needed in hdr, needed
    assert hdr == HEADERS
    dv = ws.data_validations.dataValidation[0]
    for o in OUTCOMES:
        assert o in dv.formula1


def test_decisions_round_trip_through_the_file(tmp_path):
    out = tmp_path / "r.xlsx"
    write_sheet(out, [row(), row(person_id="2", affiliation_id="11", display_name="Bob")])
    wb = load_workbook(out); ws = wb["Review"]
    col = {c.value: c.column for c in ws[1]}
    ws.cell(2, col["Actual rank"], "associate"); ws.cell(2, col["Since (year)"], 2021)
    ws.cell(2, col["Where you saw it"], "https://x"); ws.cell(3, col["Actual rank"], "unclear")
    wb.save(out)
    ds = read_decisions(out)
    assert len(ds) == 1                                    # 'unclear' is not a decision
    assert ds[0]["affiliation_id"] == "10" and ds[0]["decision"] == "associate" and ds[0]["since"] == 2021


def test_a_promotion_closes_the_row_and_opens_a_dated_one():
    affs, log = apply([decision(since=2021)], [aff()])
    old, new = affs
    assert old["end_date"] == "2021-01-01" and "closed 2021-01-01: associate" in old["source"]
    assert new["rank"] == "associate" and new["start_date"] == "2021-01-01" and new["end_date"] == ""
    assert new["appointment_type"] == "regular" and new["affiliation_id"] == 11
    assert "https://somewhere.edu/people/ada" in new["source"]


def test_no_since_year_means_the_review_date():
    affs, _ = apply([decision()], [aff()])
    assert affs[1]["start_date"] == "2026-09-17"


def test_the_same_rank_only_records_a_verification():
    affs, log = apply([decision(decision="assistant")], [aff()])
    assert len(affs) == 1 and affs[0]["end_date"] == ""
    assert "rank verified 2026-09-17" in affs[0]["source"]
    assert log[0].startswith("verified")


def test_emeritus_and_adjunct_change_the_appointment_type_and_keep_the_rank():
    affs, _ = apply([decision(decision="emeritus")], [aff(rank="associate")])
    assert affs[1]["appointment_type"] == "emeritus" and affs[1]["rank"] == "associate"
    affs, _ = apply([decision(decision="adjunct or lecturer")], [aff()])
    assert affs[1]["appointment_type"] == "adjunct" and affs[1]["rank"] == "assistant"


def test_leaving_closes_without_a_replacement():
    affs, log = apply([decision(decision="left the program")], [aff()])
    assert len(affs) == 1 and affs[0]["end_date"] == "2026-09-17"
    assert log[0].startswith("left")


def test_an_already_closed_row_is_never_touched_twice():
    affs, log = apply([decision()], [aff(end_date="2026-09-15")])
    assert len(affs) == 1 and log[0].startswith("skip")


def test_since_accepts_a_year_a_date_or_nothing():
    assert since_date(2021, "2026-09-17") == "2021-01-01"
    assert since_date("2021", "2026-09-17") == "2021-01-01"
    assert since_date("2022-08-15", "2026-09-17") == "2022-08-15"
    assert since_date("", "2026-09-17") == "2026-09-17"
    assert since_date(None, "2026-09-17") == "2026-09-17"


def test_a_rank_set_by_review_counts_as_verified_that_day():
    """Otherwise a promotion dated 2019 is re-flagged as '7 years associate'
    the moment the reviewer records it."""
    from datetime import date
    from pipeline.detect_changes import rank_review, verified_on
    affs, _ = apply([decision(since=2019)], [aff()])
    new = affs[1]
    assert verified_on(new["source"]) == date(2026, 9, 17)
    assert rank_review(new["rank"], new["start_date"], 2005, date(2026, 9, 17),
                       verified=verified_on(new["source"])) is None
