"""Regression test: the SQL views must reproduce the numbers in Tom's workbook.

The workbook stored three percentile columns, all computed by VBA macros:
  %Cites2017  PERCENTRANK of total citations over all faculty
  %Rank       PERCENTRANK of total citations within position (rank)
  %RankH      PERCENTRANK of h-index over all faculty
and a 'Current Schools' sheet with faculty count, median and mean citations
per school.

Small deviations are expected and explained:
  * Tom's arrays include the two duplicate rows (Guang Tian, Lisa Berglund);
    ours count each person once, so n differs by 2.
  * The within-rank macros used hard-coded row ranges that were one row off
    at the associate/full boundary.
Tolerances below are set just above the deviations those two effects cause.
"""

from __future__ import annotations

import csv
import json
import shutil
import sqlite3
from pathlib import Path

import openpyxl
import pytest

from pipeline import LEGACY_DIR, ROSTER_DIR, SNAPSHOT_DIR
from pipeline.build_db import build

WORKBOOK = LEGACY_DIR / "UPDATED_Cites2026.xlsm"
DUPLICATED = {"Guang Tian", "Lisa Berglund"}

# Workbook cells known to be wrong. Jacob Oluwoye's within-rank percentile is
# stored as 0.835, which is impossible for 2,430 citations among full
# professors (the correct value is 0.494). His row sits near the top of the
# sheet after the school sort, and the full-professor macro only overwrote rows
# 648 to 1052, so the cell kept a value from an earlier state of the sheet.
STALE_WITHIN_RANK_CELLS = {"Jacob Oluwoye"}

pytestmark = pytest.mark.skipif(
    not WORKBOOK.exists() or not (ROSTER_DIR / "person.csv").exists(),
    reason="handover workbook or migrated roster not present",
)


@pytest.fixture(scope="module")
def db(tmp_path_factory) -> sqlite3.Connection:
    # Build from the migrated snapshots only. Later collection runs legitimately
    # change the headline numbers, and this test is about the migration.
    snaps = tmp_path_factory.mktemp("snapshots")
    for meta in SNAPSHOT_DIR.glob("*/*.meta.json"):
        if json.loads(meta.read_text()).get("trigger") == "migration":
            dest = snaps / meta.parent.name
            dest.mkdir(exist_ok=True)
            shutil.copy(meta, dest / meta.name)
            shutil.copy(meta.with_suffix("").with_suffix(".csv"), dest / meta.with_suffix("").with_suffix(".csv").name)
    # Likewise rebuild the roster as migrated: keep only the affiliations the
    # workbook produced and undo closures applied by later change runs.
    roster = tmp_path_factory.mktemp("roster")
    for f in ROSTER_DIR.glob("*.csv"):
        shutil.copy(f, roster / f.name)
    with (ROSTER_DIR / "affiliation.csv").open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fields = reader.fieldnames
        affs = [a for a in reader if (a["source"] or "").startswith("workbook-2026")]
    for a in affs:
        if "| closed" in (a["source"] or ""):
            a["end_date"] = ""
            a["source"] = a["source"].split(" | closed")[0]
    with (roster / "affiliation.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(affs)
    path = build(tmp_path_factory.mktemp("db") / "citations.sqlite", roster, snaps, quiet=True)
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    return con


@pytest.fixture(scope="module")
def workbook():
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    cites = list(wb["Cites"].iter_rows(values_only=True))
    hdr = cites[0]
    idx = {h: i for i, h in enumerate(hdr)}
    faculty = [
        {h: r[idx[h]] for h in hdr if h}
        for r in cites[1:] if r[0]
    ]
    schools = list(wb["Current Schools"].iter_rows(values_only=True))
    shdr = schools[0]
    sidx = {h: i for i, h in enumerate(shdr) if h}
    school_rows = [
        {h: r[sidx[h]] for h in sidx}
        for r in schools[1:] if r[sidx["Planning School"]]
    ]
    return faculty, school_rows


def test_every_workbook_person_is_in_the_database(db, workbook):
    faculty, _ = workbook
    names = {f["Name"].strip() for f in faculty}
    db_names = {r[0] for r in db.execute("SELECT display_name FROM person")}
    assert names == db_names
    # 1,051 rows, two people twice -> 1,049 people
    assert len(db_names) == len(faculty) - len(DUPLICATED)


def test_percentiles_match_workbook(db, workbook):
    faculty, _ = workbook
    ours = {
        (r["display_name"], r["department"]): r
        for r in db.execute("SELECT * FROM v_person_percentiles")
    }
    worst = {"all": 0.0, "rank": 0.0, "h": 0.0}
    checked = 0
    for f in faculty:
        key = (f["Name"].strip(), f["CurrentSchool"])
        if key not in ours:
            assert f["Name"].strip() in DUPLICATED, f"missing {key}"
            continue  # the closed duplicate row has no current affiliation
        r = ours[key]
        assert r["total_citations"] == f["Cites2017"], key
        assert r["h_index"] == f["H-Index"], key
        worst["all"] = max(worst["all"], abs(r["pct_citations_all"] - f["%Cites2017"]))
        worst["h"] = max(worst["h"], abs(r["pct_h_all"] - f["%RankH"]))
        if key[0] not in STALE_WITHIN_RANK_CELLS:
            worst["rank"] = max(worst["rank"], abs(r["pct_citations_rank"] - f["%Rank"]))
        checked += 1
    assert checked == len(faculty) - len(DUPLICATED)
    print(f"\nmax |ours - workbook|: citations overall {worst['all']:.4f}, "
          f"citations within rank {worst['rank']:.4f}, h-index overall {worst['h']:.4f}")
    # Two fewer rows in a ~1,050 array moves a percentile by at most ~0.002;
    # the off-by-one rank boundary moves the within-rank value by ~0.003.
    assert worst["all"] < 0.005
    assert worst["h"] < 0.005
    assert worst["rank"] < 0.01


def test_department_summary_matches_workbook(db, workbook):
    faculty, school_rows = workbook
    ours = {r["short_name"]: r for r in db.execute("SELECT * FROM v_department_summary")}
    affected = {f["CurrentSchool"] for f in faculty if f["Name"].strip() in DUPLICATED}
    checked = 0
    for s in school_rows:
        label = s["Planning School"]
        assert label in ours, f"school {label!r} missing from v_department_summary"
        r = ours[label]
        if label in affected:
            # Tom counted the stale duplicate row here; we closed it.
            assert r["n_faculty"] in (s["Faculty"], s["Faculty"] - 1), label
            continue
        assert r["n_faculty"] == s["Faculty"], label
        # Tom's sheet stores some means to two decimals and others as whole numbers.
        assert abs(r["median_citations"] - s["Median"]) < 0.5, label
        assert abs(r["mean_citations"] - s["Mean"]) < 0.51, label
        checked += 1
    assert checked >= len(school_rows) - len(affected)


def test_headline_source_split(db):
    counts = dict(db.execute(
        "SELECT source, COUNT(*) FROM v_headline_metrics GROUP BY source"
    ).fetchall())
    assert counts["google_scholar"] == 808 - 2   # two Scholar-profile people were duplicated
    assert counts["pop"] == 243


def test_years_since_phd_has_no_division_by_zero(db):
    zero = db.execute(
        "SELECT COUNT(*) FROM v_person_percentiles WHERE years_since_phd <= 0"
    ).fetchone()[0]
    null_cpy = db.execute(
        "SELECT COUNT(*) FROM v_person_percentiles WHERE citations_per_year IS NULL"
    ).fetchone()[0]
    # Fresh PhDs have zero years; their citations_per_year must be NULL, not an error.
    assert zero == null_cpy
