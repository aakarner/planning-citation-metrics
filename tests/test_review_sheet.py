"""The review spreadsheet: banding, and the contract with import_review.

The sheet is written by one script and read back by another, so the columns
import_review looks up by name are a contract between them. A rename that broke
it would otherwise surface only after someone had spent an afternoon filling the
sheet in.
"""

from openpyxl import load_workbook

from pipeline.build_review_sheet import band_of, write_sheet


def cand(name_sim, inst_match, **kw):
    base = {"person_id": "1", "external_id": "A123", "candidate_name": "A. Person",
            "last_known_institution": "Somewhere University", "score": 0.9,
            "name_sim": name_sim, "inst_match": inst_match, "cited_by_count": 100,
            "works_count": 10, "h_index": 5}
    return {**base, **kw}


def test_bands_run_from_easiest_to_hardest():
    assert band_of(cand(1.0, 1.0)) == "A"
    assert band_of(cand(0.60, 1.0)) == "B"
    assert band_of(cand(1.0, 0.0)) == "C"
    assert band_of(cand(0.60, 0.5)) == "D"


def test_a_missing_score_does_not_raise():
    assert band_of(cand(None, None)) == "D"
    assert band_of(cand("", "")) == "D"


def _build(tmp_path, groups):
    out = tmp_path / "review.xlsx"
    write_sheet(out, groups)
    return load_workbook(out)


PERSON = {"display_name": "Ada Lovelace", "dept": "Somewhere University",
          "base_source": "pop", "base_cites": 200}


def test_the_sheet_carries_every_column_import_review_needs(tmp_path):
    wb = _build(tmp_path, [("A", PERSON, [cand(1.0, 1.0)])])
    assert "Review" in wb.sheetnames
    hdr = [c.value for c in wb["Review"][1]]
    for needed in ("Decision", "Reviewed by", "Date", "Notes", "person_id", "external_id"):
        assert needed in hdr, needed


def test_key_columns_repeat_so_the_sheet_survives_sorting(tmp_path):
    """No blank separator rows, and the person is named on every one of their
    rows: filtering or sorting a sheet with either would scramble it."""
    rows = [cand(1.0, 1.0, external_id="A1", works_count=9),
            cand(1.0, 1.0, external_id="A2", works_count=2),
            cand(1.0, 1.0, external_id="A3", works_count=1)]
    ws = _build(tmp_path, [("A", PERSON, rows)])["Review"]
    hdr = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    body = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(body) == 3
    assert all(r[col["Person"]] == "Ada Lovelace" for r in body)
    assert all(r[col["Our figure"]] == 200 for r in body)
    assert ws.auto_filter.ref == f"A1:{ws.cell(1, len(hdr)).column_letter}4"


def test_the_ratio_is_openalex_over_our_own_figure(tmp_path):
    ws = _build(tmp_path, [("A", PERSON, [cand(1.0, 1.0, cited_by_count=50)])])["Review"]
    col = {c.value: i for i, c in enumerate(ws[1])}
    assert list(ws.iter_rows(min_row=2, values_only=True))[0][col["Ratio"]] == 0.25


def test_the_decision_cells_offer_a_dropdown(tmp_path):
    ws = _build(tmp_path, [("A", PERSON, [cand(1.0, 1.0)])])["Review"]
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    assert "accepted" in dvs[0].formula1 and "rejected" in dvs[0].formula1
